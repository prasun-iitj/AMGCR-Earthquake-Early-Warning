"""One-off builder for FINAL_SUBMISSION deliverables (D-F4). Not part of analysis pipeline."""
from __future__ import annotations

import json
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "FINAL_SUBMISSION"
REPORTS = ROOT / "reports"
FIGURES = REPORTS / "figures"
DOCS = ROOT / "docs"

# Slide definitions: (title, bullets, figure_paths relative to FIGURES)
SLIDES: list[tuple[str, list[str], list[str]]] = [
    (
        "Reproducible AI-Assisted Earthquake Early Warning Research for Western Seismic Regions",
        [
            "Method Development on a California FDSN Pilot with Transfer Path to European Networks",
            "AMGCR Earthquake Research — D-F1 / D-F2",
            "Swiss certificate programme (AMGCR)",
            "29 July 2026",
        ],
        [],
    ),
    (
        "Why Earthquake Early Warning and open methods matter",
        [
            "EEW: alerts seconds to tens of seconds before strong shaking",
            "P-wave detection, magnitude estimation, ground-motion prediction",
            "European risk: Mediterranean belt, Alps, anthropogenic seismicity",
            "Switzerland/EU expect documented, open methodologies",
        ],
        [],
    ),
    (
        "Europe-focused programme; California as methods laboratory",
        [
            "Primary target: European seismic networks",
            "California IRIS/EarthScope pilot: completed validation lab",
            "EarthESND: reference only (88 tests on full checkout)",
            "Pilot science freeze: 2026-07-29",
        ],
        [],
    ),
    (
        "Gaps this workflow addresses",
        [
            "Few packaged EEW workflows for European FDSN data",
            "Missing trails: raw MiniSEED → ML-ready features",
            "Benchmark code decoupled from Western pilot science",
            "Small-N pilots cannot support operational EEW claims",
        ],
        [],
    ),
    (
        "Objectives and completion status",
        [
            "Primary: reproducible, AI-ready workflow for Europe",
            "O1–O6 completed (through 8×18 features + interpretation)",
            "O4: instrument response 0/8 applied — documented",
            "O7–O9 proposed: Europe, ML, real-time",
        ],
        [],
    ),
    (
        "Literature — onsite EEW",
        [
            "STA/LTA and triggers for P-wave onset",
            "Allen & Kanamori (2003); Hoshiba et al. (2008)",
            "Trade-offs: latency, false alarms, magnitude saturation",
            "European focus: metadata, latency budgets, uncertainty",
        ],
        [],
    ),
    (
        "Literature — ObsPy, FDSN, ML, reproducibility",
        [
            "ObsPy (Krischer et al., 2015)",
            "USGS events + EarthScope waveforms in pilot",
            "EarthESND: ESN/DENN; reference code only here",
            "Manifests, JSON configs, phase reports",
        ],
        [],
    ),
    (
        "Dataset — completed California pilot",
        [
            "8 events; M 4.05–4.87; BHZ; 40 Hz; 300 s from origin",
            "CI.ADO (6) · CI.USC (2)",
            "12 MiniSEED on disk at EDA; 8 manifest-linked analyses",
            "No European waveforms in completed pilot",
        ],
        [],
    ),
    (
        "How data were acquired",
        [
            "USGS: California 2024, M 4.0–7.5 → 41 catalog matches",
            "EarthScope: CI, NC, BK; 8 events retained",
            "Distance-ranked station list — not uniform sampling",
            "Public FDSN; no credentials",
        ],
        [],
    ),
    (
        "California pilot workflow",
        [
            "Done: download → EDA → signal → preprocessing → features",
            "Done: interpretation, D-S1/D-S2, D-F1 report",
            "EarthESND src/ not used for pilot results",
            "Proposed: EU FDSN, larger N, ML, SeedLink",
        ],
        [],
    ),
    (
        "Methods — EDA & signal (Phase A)",
        [
            "EDA: ObsPy inspection; four cohort figures",
            "Signal: peak/RMS, 0–5 s noise proxy, peak SNR",
            "STA/LTA: 0.5 s / 10 s; threshold 2.5",
            "Windows start at origin — no pre-event quiet time",
        ],
        [],
    ),
    (
        "Methods — preprocessing & features (Phase B)",
        [
            "Demean, detrend, 5% Hann; band-pass 0.1–15 Hz",
            "Response removal attempted → 0/8 applied",
            "18 features from filtered traces",
            "PGM/Arias: counts-based proxies only",
        ],
        [],
    ),
    (
        "Results — EDA (example waveform)",
        [
            "Figure 1: M 4.87 at CI.ADO — raw counts",
            "Narrow M 4–5 band; query cutoffs",
            "300 s window from trace start",
        ],
        ["example_waveform_raw.png"],
    ),
    (
        "Results — EDA (catalog views)",
        [
            "Figures 2–4: magnitude, time coverage, station usage",
            "ADO supplies 6/8 traces — acquisition bias",
            "Gaps reflect download success",
        ],
        ["magnitude_histogram.png", "events_over_time.png", "station_usage_frequency.png"],
    ),
    (
        "Results — signal analysis (ci40675215)",
        [
            "P pick: 10.15–55.15 s (mean 31.3 s)",
            "Peak SNR: 6.3–778 (median 49.9)",
            "Figure 5: high-SNR case, early pick ~10.2 s",
        ],
        ["signal_analysis_ci40675215.png"],
    ),
    (
        "Results — signal analysis (ci40699207 / USC)",
        [
            "ADO: mean SNR 169; USC: mean SNR 42",
            "Figure 6: SNR ~6; duration 0 s",
            "SNR is QC indicator, not detection probability",
        ],
        ["signal_analysis_ci40699207.png"],
    ),
    (
        "Results — preprocessing",
        [
            "Mean peak-SNR: 137 → 1684 (filtered)",
            "Figures 7–8: raw → detrended → filtered → normalised",
            "Filtered counts ≠ physical motion without response",
        ],
        ["preprocessing_ci40675215.png", "preprocessing_ci40699207.png"],
    ),
    (
        "Results — features (spectra & correlations)",
        [
            "8×18 feature matrix",
            "Crest factor mean ≈ 11.4; centroid ~2.3–2.6 Hz",
            "Figures 9–10: FFT grid, correlation heatmap",
            "No ML training in completed pilot",
        ],
        ["feature_engineering_fft_spectra.png", "feature_correlation_heatmap.png"],
    ),
    (
        "Results — feature distributions & stations",
        [
            "Figures 11–12: histograms, ADO vs USC boxplots",
            "Amplitude features: high collinearity",
            "USC n=2 — indicative only",
        ],
        ["feature_distributions.png", "feature_boxplots_by_station.png"],
    ),
    (
        "Discussion",
        [
            "Script-driven FDSN → features + NPZ stages",
            "Late picks + origin noise limit trigger tuning",
            "Feature matrix: exploratory baselines only",
            "EarthESND: future comparison, not validated here",
        ],
        [],
    ),
    (
        "Limitations",
        [
            "N=8; single BHZ; origin-aligned windows",
            "0/8 response removal; ADO-heavy sampling",
            "Mixed mw/ml; unvalidated P picks",
            "Not comparable to EarthESND K-NET metrics",
        ],
        [],
    ),
    (
        "Threats to validity",
        [
            "Construct: counts vs physical units",
            "External: California ≠ Europe → EIDA/ORFEUS pilot",
            "Statistical: N=8 → larger catalogues",
            "Mitigations are proposed future work",
        ],
        [],
    ),
    (
        "European context & transfer",
        [
            "Europe primary focus; California validates methods",
            "ObsPy/FDSN reproducibility demonstrated",
            "Proposed: ORFEUS/EIDA manifests + same scripts",
            "No EU waveform data yet",
        ],
        [],
    ),
    (
        "Future work, conclusions & thank you",
        [
            "Proposed: 50–200+ events, EU pilot, response fix, ML, SeedLink",
            "Conclusion: reproducible chain to 8×18 features established",
            "Acknowledgements: USGS, EarthScope, ObsPy",
            "Questions?",
        ],
        [],
    ),
]

CAPTIONS = {
    "example_waveform_raw.png": "Figure 1. Raw BHZ, M 4.87, CI.ADO (digital counts).",
    "magnitude_histogram.png": "Figure 2. Magnitude distribution (N=8).",
    "events_over_time.png": "Figure 3. Origin time vs magnitude.",
    "station_usage_frequency.png": "Figure 4. Station usage (ADO vs USC).",
    "signal_analysis_ci40675215.png": "Figure 5. Signal analysis — ci40675215.",
    "signal_analysis_ci40699207.png": "Figure 6. Signal analysis — ci40699207 (USC).",
    "preprocessing_ci40675215.png": "Figure 7. Preprocessing — ci40675215.",
    "preprocessing_ci40699207.png": "Figure 8. Preprocessing — ci40699207.",
    "feature_engineering_fft_spectra.png": "Figure 9. FFT power spectra.",
    "feature_correlation_heatmap.png": "Figure 10. Feature correlation heatmap.",
    "feature_distributions.png": "Figure 11. Feature distributions.",
    "feature_boxplots_by_station.png": "Figure 12. Features by station.",
}


def run_pandoc(md: Path, docx: Path, *, toc: bool = False, resource_path: Path | None = None) -> None:
    cmd = [
        "pandoc",
        str(md),
        "-f",
        "markdown",
        "-t",
        "docx",
        "-o",
        str(docx),
    ]
    if toc:
        cmd.append("--toc")
        cmd.extend(["--toc-depth", "3"])
    if resource_path:
        cmd.extend(["--resource-path", str(resource_path)])
    subprocess.run(cmd, check=True, cwd=md.parent)


def docx_to_pdf(docx: Path, pdf: Path) -> None:
    from docx2pdf import convert

    convert(str(docx), str(pdf))


def build_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for idx, (title, bullets, figs) in enumerate(SLIDES, start=1):
        slide = prs.slides.add_slide(blank)
        # White background
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.35), Inches(12.3), Inches(1.0))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Calibri"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

        top = 1.35
        body_h = 2.0 if figs else 5.0
        body_box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.2), Inches(body_h))
        btf = body_box.text_frame
        btf.word_wrap = True
        for i, line in enumerate(bullets[:6]):
            para = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
            para.text = line
            para.level = 0
            para.font.name = "Calibri"
            para.font.size = Pt(20)
            para.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        if figs:
            n = len(figs)
            y0 = top + body_h + 0.15
            usable_w = 12.2
            if n == 1:
                positions = [(0.55, y0, usable_w, 3.6)]
            elif n == 2:
                w = (usable_w - 0.3) / 2
                positions = [(0.55, y0, w, 3.2), (0.55 + w + 0.3, y0, w, 3.2)]
            else:
                w = (usable_w - 0.4) / 3
                positions = [
                    (0.55, y0, w, 2.5),
                    (0.55 + w + 0.2, y0, w, 2.5),
                    (0.55 + 2 * (w + 0.2), y0, w, 2.5),
                ]
            for (fname, (left, y, w, h)) in zip(figs, positions):
                img = FIGURES / fname
                if img.is_file():
                    slide.shapes.add_picture(str(img), Inches(left), Inches(y), width=Inches(w))
                    cap = slide.shapes.add_textbox(Inches(left), Inches(y + h + 0.05), Inches(w), Inches(0.35))
                    cp = cap.text_frame.paragraphs[0]
                    cp.text = CAPTIONS.get(fname, fname)
                    cp.font.name = "Calibri"
                    cp.font.size = Pt(12)
                    cp.font.italic = True
                    cp.alignment = PP_ALIGN.CENTER

        num = slide.shapes.add_textbox(Inches(12.5), Inches(7.05), Inches(0.6), Inches(0.3))
        np = num.text_frame.paragraphs[0]
        np.text = str(idx)
        np.font.name = "Calibri"
        np.font.size = Pt(14)
        np.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        np.alignment = PP_ALIGN.RIGHT

    prs.save(path)


def pptx_to_pdf(pptx: Path, pdf: Path) -> None:
    try:
        import comtypes.client  # type: ignore

        powerpoint = comtypes.client.CreateObject("Powerpoint.Application")
        powerpoint.Visible = 1
        deck = powerpoint.Presentations.Open(str(pptx.resolve()), WithWindow=False)
        deck.SaveAs(str(pdf.resolve()), 32)  # ppSaveAsPDF
        deck.Close()
        powerpoint.Quit()
    except Exception as exc:
        print(f"WARNING: PowerPoint PDF export failed: {exc}", file=sys.stderr)


def validate() -> dict:
    expected_figs = set()
    for _, _, figs in SLIDES:
        expected_figs.update(figs)
    missing_figs = [f for f in expected_figs if not (FIGURES / f).is_file()]

    outputs = {
        "Research_Report_Final.docx": OUT / "Research_Report_Final.docx",
        "Research_Report_Final.pdf": OUT / "Research_Report_Final.pdf",
        "Presentation.pptx": OUT / "Presentation.pptx",
        "Presentation.pdf": OUT / "Presentation.pdf",
        "Presentation_Script.docx": OUT / "Presentation_Script.docx",
        "Presentation_Script.pdf": OUT / "Presentation_Script.pdf",
        "REPRODUCIBILITY_STATEMENT.pdf": OUT / "REPRODUCIBILITY_STATEMENT.pdf",
        "DATA_AVAILABILITY_STATEMENT.pdf": OUT / "DATA_AVAILABILITY_STATEMENT.pdf",
    }
    missing_files = [k for k, p in outputs.items() if not p.is_file()]

    return {
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "figures_expected_in_deck": sorted(expected_figs),
        "figures_missing_on_disk": missing_figs,
        "outputs_missing": missing_files,
        "figure_count_on_disk": len(list(FIGURES.glob("*.png"))) if FIGURES.is_dir() else 0,
    }


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)

    run_pandoc(
        REPORTS / "Research_Report_Final.md",
        OUT / "Research_Report_Final.docx",
        toc=True,
        resource_path=REPORTS,
    )
    run_pandoc(
        REPORTS / "Presentation_Outline_D-F2.md",
        OUT / "Presentation_Outline_D-F2.docx",
        resource_path=REPORTS,
    )
    run_pandoc(REPORTS / "Presentation_Script_D-F3.md", OUT / "Presentation_Script.docx")

    run_pandoc(DOCS / "REPRODUCIBILITY_STATEMENT.md", OUT / "_REPRODUCIBILITY.docx", resource_path=ROOT)
    run_pandoc(DOCS / "DATA_AVAILABILITY_STATEMENT.md", OUT / "_DATA_AVAILABILITY.docx", resource_path=ROOT)

    for docx_name, pdf_name in [
        ("Research_Report_Final.docx", "Research_Report_Final.pdf"),
        ("Presentation_Script.docx", "Presentation_Script.pdf"),
        ("_REPRODUCIBILITY.docx", "REPRODUCIBILITY_STATEMENT.pdf"),
        ("_DATA_AVAILABILITY.docx", "DATA_AVAILABILITY_STATEMENT.pdf"),
    ]:
        docx_to_pdf(OUT / docx_name, OUT / pdf_name)

    build_pptx(OUT / "Presentation.pptx")
    pptx_to_pdf(OUT / "Presentation.pptx", OUT / "Presentation.pdf")

    (OUT / "_REPRODUCIBILITY.docx").unlink(missing_ok=True)
    (OUT / "_DATA_AVAILABILITY.docx").unlink(missing_ok=True)

    result = validate()
    (OUT / "VALIDATION_SUMMARY.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    status = "PASS" if not result["figures_missing_on_disk"] and not result["outputs_missing"] else "FAIL"
    md = [
        "# Final submission validation summary",
        "",
        f"**Overall:** **{status}**",
        f"**UTC:** {result['timestamp_utc']}",
        "",
        "## Figures",
        f"- PNG count on disk (`reports/figures/`): **{result['figure_count_on_disk']}**",
        f"- Figures required in deck: **{len(result['figures_expected_in_deck'])}**",
        f"- Missing on disk: **{len(result['figures_missing_on_disk'])}** "
        + (f"`{result['figures_missing_on_disk']}`" if result["figures_missing_on_disk"] else "(none)"),
        "",
        "## Required outputs",
    ]
    for name in [
        "Research_Report_Final.docx",
        "Research_Report_Final.pdf",
        "Presentation.pptx",
        "Presentation.pdf",
        "Presentation_Script.docx",
        "Presentation_Script.pdf",
        "REPRODUCIBILITY_STATEMENT.pdf",
        "DATA_AVAILABILITY_STATEMENT.pdf",
    ]:
        ok = (OUT / name).is_file()
        md.append(f"- {'[x]' if ok else '[ ]'} `{name}` — **{'PASS' if ok else 'FAIL'}**")
    md.append("")
    md.append("## Notes")
    md.append("- PDF generation requires Microsoft Word (docx2pdf) and PowerPoint (presentation PDF).")
    md.append("- Re-run: `python FINAL_SUBMISSION/_build_submission_assets.py`")
    (OUT / "VALIDATION_SUMMARY.md").write_text("\n".join(md), encoding="utf-8")
    print(json.dumps(result, indent=2))
    return 0 if status == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
